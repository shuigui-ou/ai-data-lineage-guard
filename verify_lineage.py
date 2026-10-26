#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verify_lineage.py — 数值血缘对账器（防 AI 偷懒核心工具）

读入「原始数据（真值）」与「AI 产出的带血缘标签结果」，机械地逐一核对，
抓出：抄错/算错（值≠源单元格）、派生统计不符、以及来源缺失（懒得查）的红色信号。

核心仅用 Python 标准库；xlsx / parquet 为可选依赖（缺则显式报错 exit 3，绝不降级）。数据源支持 .csv / .json / .xlsx / .parquet，产出支持 .json 清单 或 .md/.txt 行内标签。

依赖分级（缺依赖时显式报错，绝不降级、绝不静默跳过）：
    · csv/json  → 纯标准库，任何环境可跑
    · xlsx      → 需要 openpyxl（缺则 exit 3，提示 pip install openpyxl）
    · parquet   → 需要 pyarrow（缺则 exit 3，提示 pip install pyarrow）
"""
import sys
import json
import csv
import re
import datetime
import argparse
from pathlib import Path

# 数据源描述（供报告打印，含 xlsx 的 sheet 名、parquet 的行列数等元信息）
LAST_SOURCE_DESC = ""


# ───────────────────────── 数据源加载 ─────────────────────────

def load_source(path):
    """返回 (header: list[str], rows: list[dict])。按扩展名分派；xlsx 走 openpyxl 只读分支。"""
    global LAST_SOURCE_DESC
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _load_xlsx(p)
    if suffix in (".parquet", ".pq"):
        return _load_parquet(p)
    if suffix == ".json":
        with open(p, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            # 容许多包一层：找第一个值是 list 的键
            for v in data.values():
                if isinstance(v, list) and v and isinstance(v[0], dict):
                    data = v
                    break
        header = list(data[0].keys())
        rows = [{(k): ("" if v is None else str(v)) for k, v in row.items()} for row in data]
        LAST_SOURCE_DESC = str(p)
        return header, rows
    # 默认 csv
    with open(p, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        header = reader.fieldnames or []
        rows = [dict(r) for r in reader]
    LAST_SOURCE_DESC = str(p)
    return header, rows


def _load_xlsx(p):
    """xlsx 源：openpyxl 只读 + data_only（取缓存计算值）。

    五个坑的处理：
    1. read_only=True 省内存；data_only=True 取公式的缓存计算值而非公式串。
    2. 公式无缓存值（文件未经 Excel 保存）→ 显式报错，绝不当空值静默过（防假 PASS）。
    3. 首行当 header，R1=首条数据行，与 csv 分支完全一致；列按表头名引用（同 csv）。
    4. 默认取第一个 sheet，并在报告标注 sheet 名，防"读错表还 PASS"。
    5. 缺 openpyxl 依赖 → 显式 ImportError，不降级、不跳过。
    """
    global LAST_SOURCE_DESC
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "XLSX 源需要 openpyxl，请 pip install openpyxl"
            "（缺依赖时显式报错，不降级不跳过）"
        )

    # 第一遍：data_only=True 取缓存计算值
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]  # 默认第一个 sheet
    sheet_name = ws.title
    none_cells = set()
    raw = []
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        raw.append(row)
        for c_idx, val in enumerate(row):
            if val is None:
                none_cells.add((r_idx, c_idx))
    wb.close()

    # 第二遍：data_only=False 检测"公式单元格无缓存值"
    wb2 = openpyxl.load_workbook(str(p), read_only=True, data_only=False)
    ws2 = wb2[wb2.sheetnames[0]]
    formula_no_cache = 0
    for r_idx, row in enumerate(ws2.iter_rows(values_only=True), start=1):
        for c_idx, val in enumerate(row):
            if isinstance(val, str) and val.startswith("=") and (r_idx, c_idx) in none_cells:
                formula_no_cache += 1
    wb2.close()

    if formula_no_cache > 0:
        raise ValueError(
            f"xlsx 含 {formula_no_cache} 个公式单元格无缓存计算值"
            f"（文件可能未经 Excel 保存，data_only 读到 None）。无法验证："
            f"请用 Excel 打开并保存一次，或将公式转为值后再校验。"
        )

    if not raw:
        raise ValueError("xlsx 空表：无任何行")
    header_row = raw[0]
    header = [(str(h).strip() if h is not None else f"col{i + 1}") for i, h in enumerate(header_row)]
    rows = []
    for data_row in raw[1:]:
        d = {}
        all_none = True
        for i, h in enumerate(header):
            val = data_row[i] if i < len(data_row) else None
            if val is not None:
                all_none = False
            d[h] = "" if val is None else (val if isinstance(val, (int, float, bool)) else str(val))
        if all_none:
            continue
        rows.append(d)
    LAST_SOURCE_DESC = f"{p} (sheet: {sheet_name})"
    return header, rows


def _parquet_cell(v):
    """parquet 单元格 → 与 csv/json/xlsx 分支一致的字符串表示。

    归一化规则（与现有比较逻辑对齐）：
      None / NaN        → ""        （空值不当 "nan" 文本参与字符串身份比对）
      bool              → "True"/"False"
      datetime/date/time→ ISO 8601 字符串
      Decimal / 数值    → str()
      bytes             → utf-8 解码（失败则 replace）
      list/dict（嵌套列）→ JSON 文本
    """
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, float) and v != v:  # NaN
        return ""
    if isinstance(v, (bytes, bytearray, memoryview)):
        return bytes(v).decode("utf-8", errors="replace")
    if isinstance(v, datetime.datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.isoformat()
    if isinstance(v, (list, tuple, dict)):
        try:
            return json.dumps(v, ensure_ascii=False, default=str)
        except (TypeError, ValueError):
            return str(v)
    return str(v)


def _load_parquet(p):
    """parquet 源：pyarrow 分批读取（可选依赖）。

    四个坑的处理：
    1. pyarrow 是**可选依赖**，缺失 → 显式 ImportError（exit 3），绝不降级/静默跳过。
    2. 分批读（iter_batches）而非全量 to_pylist，避免大表打爆内存。
    3. 列名取 schema（而不是首行），空表也报"无任何行"，不静默 PASS。
    4. 值统一转字符串（NaN/None → ""），与 json 分支一致，文本身份比对不失真。
    """
    global LAST_SOURCE_DESC
    try:
        import pyarrow.parquet as pq
    except ImportError:
        raise ImportError(
            "Parquet 源需要 pyarrow，请 pip install pyarrow"
            "（缺依赖时显式报错，不降级不跳过）"
        )

    try:
        pf = pq.ParquetFile(str(p))
    except Exception as e:  # 损坏文件/非 parquet → 干净报错，不漏 traceback
        raise ValueError(f"parquet 无法打开（文件损坏或不是 parquet）: {e}")

    header = list(pf.schema_arrow.names)
    if not header:
        raise ValueError("parquet 空表：无任何列")

    rows = []
    empty_rows = 0
    try:
        for batch in pf.iter_batches(batch_size=4096):
            for d in batch.to_pylist():
                row = {h: _parquet_cell(d.get(h)) for h in header}
                if all(v == "" for v in row.values()):
                    empty_rows += 1
                # v0.5.2 修正 F：全空行**保留**。此前静默 continue 丢弃会让
                # source.row 与物理行整体错位，而每条 claim 仍可能 PASS ——
                # 那正是本工具最该防的"实体错位"，却成了隐性盲区。
                rows.append(row)
    except Exception as e:
        raise ValueError(f"parquet 读取失败: {e}")

    if not rows:
        raise ValueError("parquet 空表：无任何行")

    n_rows_file = pf.metadata.num_rows
    align = ("物理行对齐" if len(rows) == n_rows_file
             else f"⚠行号与物理行不一致(文件 {n_rows_file} 行/读取 {len(rows)} 行)")
    empty_note = f", 含全空行 {empty_rows} 条(已保留以对齐)" if empty_rows else ""
    LAST_SOURCE_DESC = (f"{p} (parquet: {len(header)} 列 x {n_rows_file} 行, "
                        f"读取 {len(rows)} 行{empty_note}; {align})")
    return header, rows


# ───────────────────────── 数值与派生工具 ─────────────────────────

def coerce_num(s):
    if isinstance(s, (int, float)):
        return float(s)
    if s is None:
        return None
    s = str(s).replace(",", "").replace("%", "").replace("¥", "").replace("$", "").strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _looks_numeric(s):
    """纯数字文本（含千分位/货币/百分号）走数值比对，其余走文本身份比对。"""
    return coerce_num(s) is not None


def cell_value(header, rows, row, col):
    """row: 1-based 数据行索引（不含表头）。"""
    if row < 1 or row > len(rows):
        raise IndexError(f"行越界: {row} (共 {len(rows)} 数据行)")
    if col not in header:
        raise KeyError(f"列不存在: {col} (可用列: {header})")
    return rows[row - 1][col]


def _resolve_rows(rows_spec, n_rows):
    """把 rows 声明解析成 1-based 行号索引序列。

    支持形态（v0.5.2）：
      - None / ""                  → 全部数据行
      - []                         → **空子集（0 行）**（v0.5.2 修正 E：
                                     此前 [] 被误当"全部行"，导致空子集无法表达）
      - "1-100" / "1 - 100"        → 连续闭区间
      - "41-829,1012-1034"         → 多区间并集（v0.5.2 新增；不连续行段无需逐个列举）
      - [1,4,6,11] / "1,4,6,11"    → 离散行集
    离散行集是显式列举，不做任何"就近取整/补全"推断——标签写什么就核什么。
    """
    # None / 空串 = 全部行；空列表 = 空集（两者语义不同，不能合并判断）
    if rows_spec is None:
        return list(range(1, n_rows + 1))
    if isinstance(rows_spec, str) and rows_spec.strip() == "":
        return list(range(1, n_rows + 1))

    # 离散行集：list/tuple（空列表 → 空集）
    if isinstance(rows_spec, (list, tuple)):
        if len(rows_spec) == 0:
            return []
        idxs = []
        for x in rows_spec:
            try:
                idxs.append(int(x))
            except (TypeError, ValueError):
                raise ValueError(f"离散行集含非法行号: {x!r}")
        return idxs

    s = str(rows_spec).strip()
    parts = [p.strip() for p in s.split(",") if p.strip()]
    if not parts:
        return []

    # 多段（逗号）：每段可为 "a-b" 区间或单个行号 → 并集
    if len(parts) > 1:
        idxs = []
        for part in parts:
            m = re.match(r"^(\d+)\s*-\s*(\d+)$", part)
            if m:
                a, b = int(m.group(1)), int(m.group(2))
                if a > b:
                    raise ValueError(f"rows 区间首尾颠倒: {a}-{b}")
                idxs.extend(range(a, b + 1))
            else:
                try:
                    idxs.append(int(part))
                except ValueError:
                    raise ValueError(f"离散行集含非法行号: {part!r}")
        return idxs

    m = re.match(r"^(\d+)\s*-\s*(\d+)$", s)
    if not m:
        raise ValueError(f"无法解析 rows 声明: {rows_spec!r}"
                         f"（支持 '1-100' / '41-829,1012-1034' / [1,4,6] / '1,4,6'）")
    a, b = int(m.group(1)), int(m.group(2))
    if a > b:
        raise ValueError(f"rows 区间首尾颠倒: {a}-{b}")
    return list(range(a, b + 1))


def _is_null_cell(s):
    """单元格是否为空值（v0.5.2 新增，供 nulls 算子用）。

    只认"真空"：None / 空串 / 字面 None/null/NaN。
    **不把"非数字文本"当空值** —— 否则对字符串列跑 nulls 会算出满列皆空，
    语义完全错（"github" 不是空值，只是不是数字）。
    """
    if s is None:
        return True
    t = str(s).strip()
    return t == "" or t.lower() in ("none", "null", "nan")


def recompute(header, rows, op, col, rows_spec=None):
    """对列做聚合。

    算子：
      sum / mean / count / min / max —— 只统计**可转数值**的值（count 数的是可转数值的个数）
      countrows                        —— 数**行数**（与列类型无关，字符串列也适用）
      nulls                            —— 数**空值个数**（v0.5.2 新增；空串/None/null/NaN，
                                          不把非数字文本算作空）
      countdistinct                    —— 数**去重后的取值个数**（v0.5.2 新增；空值不计入；
                                          数值按值归一，100 与 100.0 视为同一个）
    rows：None / "" / [] / "1-100" / "41-829,1012-1034" / [1,4,6] / "1,4,6"（见 _resolve_rows）
    """
    if op == "countrows":
        # 行数算子：不要求 col 存在，但若给了 col 且不存在则显式报错（防笔误被静默放过）
        if col and col not in header:
            raise KeyError(f"列不存在: {col}")
        idxs = _resolve_rows(rows_spec, len(rows))
        for i in idxs:
            if i < 1 or i > len(rows):
                raise IndexError(f"行越界: {i} (共 {len(rows)} 数据行)")
        return float(len(idxs))

    if col not in header:
        raise KeyError(f"列不存在: {col}")
    idxs = _resolve_rows(rows_spec, len(rows))
    for i in idxs:
        if i < 1 or i > len(rows):
            raise IndexError(f"行越界: {i} (共 {len(rows)} 数据行)")

    if op == "nulls":
        return float(sum(1 for i in idxs if _is_null_cell(rows[i - 1][col])))

    if op == "countdistinct":
        seen = set()
        for i in idxs:
            raw = rows[i - 1][col]
            if _is_null_cell(raw):
                continue
            v = coerce_num(raw)
            seen.add(("n", v) if v is not None else ("s", str(raw)))
        return float(len(seen))

    vals = []
    for i in idxs:
        v = coerce_num(rows[i - 1][col])
        if v is not None:
            vals.append(v)
    if op == "sum":
        return sum(vals)
    if op == "mean":
        return sum(vals) / len(vals) if vals else 0.0
    if op == "count":
        return float(len(vals))
    if op == "min":
        return min(vals) if vals else 0.0
    if op == "max":
        return max(vals) if vals else 0.0
    raise ValueError(f"不支持的算子: {op}（支持 sum/mean/count/min/max/countrows/nulls/countdistinct）")


# ───────────────────────── 单条核对 ─────────────────────────

def check_claim(claim, header, rows):
    label = claim.get("label", "(未命名)")
    value = claim.get("value")
    src = claim.get("source")
    der = claim.get("derived")

    if src:
        try:
            actual = cell_value(header, rows, int(src.get("row")), src.get("col"))
            # 字符串身份校验：源是文本时必须逐字符比对，不能 coerce_num 转数值
            if isinstance(actual, str) and not _looks_numeric(actual):
                if str(value) == actual:
                    return ("ok", label, f"源[{src.get('col')} 行{src.get('row')}] 文本匹配 = {actual!r}")
                return ("fail", label,
                        f"文本错位: 报告={value!r} 实际源={actual!r} ([{src.get('col')} 行{src.get('row')}])")
            av, bv = coerce_num(actual), coerce_num(value)
            if av is None or bv is None:
                return ("warn", label, f"数值无法解析: 报告={value} 源={actual}")
            if abs(av - bv) < 1e-6:
                return ("ok", label, f"源[{src.get('col')} 行{src.get('row')}] = {av}")
            return ("fail", label, f"抄错: 报告={bv} 实际源={av} ([{src.get('col')} 行{src.get('row')}])")
        except (IndexError, KeyError) as e:
            return ("fail", label, f"来源无效: {e}")

    if der:
        try:
            actual = recompute(header, rows, der.get("op"), der.get("col"), der.get("rows"))
            bv = coerce_num(value)
            if bv is None:
                return ("warn", label, f"报告值无法解析: {value}")
            if abs(actual - bv) < 1e-6:
                return ("ok", label, f"{der.get('op')}({der.get('col')}) 重算 = {actual}")
            return ("fail", label, f"派生不符: 报告={bv} 重算={actual} ({der.get('op')} {der.get('col')})")
        except (KeyError, ValueError, IndexError) as e:
            return ("fail", label, f"派生无效: {e}")

    return ("warn", label, "来源缺失：未标注 provenance（红色信号）")


# ───────────────────────── 解析 AI 产出 ─────────────────────────

def load_claims(path):
    p = Path(path)
    if p.suffix.lower() == ".json":
        with open(p, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            return data.get("claims", [])
        if isinstance(data, list):
            return data
        raise ValueError("JSON 产出必须是 {claims:[...]} 或 [...]")
    # 文本：解析行内标签
    with open(p, encoding="utf-8") as f:
        return parse_text_claims(f.read())


def check_claim_ids(claims):
    """claim id 唯一性检查（结论引用链的地基）。

    修正 3：id 必须唯一——否则"结论引用了哪个"有歧义，
    结论可以挑对自己有利的那个绕过检查。
    返回 (ok: bool, dup_ids: list)
    """
    seen = {}
    dups = []
    for c in claims:
        cid = c.get("id")
        if cid is None:
            continue          # id 是可选的（向后兼容）
        key = str(cid)
        if key in seen:
            dups.append(key)
        else:
            seen[key] = True
    return (len(dups) == 0), dups


def parse_text_claims(text):
    claims = []
    for line in text.splitlines():
        line = line.strip()
        if not line or "=" not in line:
            continue
        m = re.match(r"^(.*?)=\s*([-0-9.,%]+)", line)
        if not m:
            continue
        label, value = m.group(1).strip(), m.group(2).strip()
        claim = {"label": label, "value": value}
        src_m = re.search(r"来源[:：]\s*(?:文件\s*\S+\s*)?行\s*(\d+)\s*列\s*(\S+)", line)
        if src_m:
            claim["source"] = {"row": int(src_m.group(1)), "col": src_m.group(2)}
        der_m = re.search(r"派生[:：]\s*(\w+)\s*\(\s*(\S+?)\s*(?:,\s*行\s*([\d,\-]+))?\s*\)", line)
        if der_m:
            claim["derived"] = {"op": der_m.group(1), "col": der_m.group(2), "rows": der_m.group(3)}
        claims.append(claim)
    return claims


# ───────────────────────── 报告 ─────────────────────────

ICON = {"ok": "✅", "fail": "❌", "warn": "⚠️"}

def main():
    ap = argparse.ArgumentParser(description="数值血缘对账器")
    ap.add_argument("--source", required=True, help="原始数据文件 (.csv/.json/.xlsx/.parquet)")
    ap.add_argument("--claims", required=True, help="AI 产出 (.json 清单 或 .md/.txt 行内标签)")
    args = ap.parse_args()

    try:
        header, rows = load_source(args.source)
    except ImportError as e:
        print(f"❌ 依赖缺失: {e}")
        sys.exit(3)
    except FileNotFoundError:
        print(f"❌ 原始数据文件不存在: {args.source}")
        sys.exit(3)
    except (ValueError, KeyError, IndexError, OSError) as e:
        print(f"❌ 原始数据无法加载: {e}")
        sys.exit(3)

    claims = load_claims(args.claims)
    if not claims:
        print("⚠️ 未解析到任何带标签的数值声明，请检查产出格式。")
        sys.exit(2)

    # claim id 唯一性（结论引用链的地基；id 可选，故仅在有 id 时生效）
    ids_ok, dups = check_claim_ids(claims)
    if not ids_ok:
        print(f"❌ claim id 重复: {sorted(set(dups))}（id 必须唯一，否则结论引用有歧义）")
        # v0.5.2 修正 I：这是**产出内容缺陷**，不是环境/依赖问题。
        # 此前误用 exit 3（"源加载失败或依赖缺失"），调用方会以为可重试，
        # 实际必须改 claims 才能过 —— 语义上应为 FAIL(1)。
        sys.exit(1)

    results = [check_claim(c, header, rows) for c in claims]

    print("=" * 64)
    print("数值血缘对账报告")
    print(f"原始数据: {LAST_SOURCE_DESC}  ({len(rows)} 数据行, 列: {header})")
    print(f"待核声明: {len(results)} 条")
    print("=" * 64)
    for status, label, detail in results:
        print(f"{ICON[status]} {label}: {detail}")
    print("-" * 64)

    n_ok = sum(1 for r in results if r[0] == "ok")
    n_fail = sum(1 for r in results if r[0] == "fail")
    n_warn = sum(1 for r in results if r[0] == "warn")
    print(f"命中 {n_ok} ｜ 错位/不符 {n_fail} ｜ 无法验证 {n_warn}")
    verdict = "PASS" if n_fail == 0 and n_warn == 0 else "FAIL"
    print(f"结论: {verdict}")
    print("=" * 64)

    sys.exit(0 if verdict == "PASS" else 1)


if __name__ == "__main__":
    main()
